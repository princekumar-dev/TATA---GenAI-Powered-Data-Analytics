# %% [markdown]
# # Task 3 — Business Summary Report (Tata iQ Forage Simulation)
#
# Fills the official `Updated_Business_Summary_Report_Template.docx` with evidence-grounded content
# from Task 1 (EDA) and Task 2 (model plan). All statistics come from the executed Task 1 analysis
# and the executed Task 2 baseline check (`Task_2_ModelPlan/baseline_check_results.json`).
#
# Input : ../Updated_Business_Summary_Report_Template.docx  (preserved, never modified)
# Output: Task_3_Business_Report.docx  (target: <= 2 pages)

# %%
import json
from pathlib import Path

import docx
from docx.shared import Pt
from docx.enum.text import WD_LINE_SPACING

BASE = Path(__file__).resolve().parent
ROOT = BASE.parent
TEMPLATE = ROOT / "Updated_Business_Summary_Report_Template.docx"
OUT = BASE / "Task_3_Business_Report.docx"

baseline = json.loads((ROOT / "Task_2_ModelPlan" / "baseline_check_results.json").read_text(encoding="utf-8"))
pr_auc_lr = baseline["models"]["Logistic Regression (balanced)"]["pr_auc_mean"]
pr_auc_gb = baseline["models"]["HistGradientBoosting"]["pr_auc_mean"]
pr_auc_dm = baseline["models"]["Dummy (stratified)"]["pr_auc_mean"]
base_rate = baseline["base_rate"]

doc = docx.Document(TEMPLATE)
paras = doc.paragraphs

def clear_runs(p):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)

def set_text(p, text, size=None, bold=False):
    clear_runs(p)
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)

def set_labeled(p, label, content, size=9.5):
    clear_runs(p)
    r1 = p.add_run(label); r1.bold = True; r1.font.size = Pt(size)
    r2 = p.add_run(content); r2.font.size = Pt(size)

def add_par_after(anchor, text, style="Normal", bold=False, size=9.5):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text); r.bold = bold
    if size:
        r.font.size = Pt(size)
    anchor._p.addnext(p._p)
    return p

# ---------------- 1. Summary of Predictive Insights ----------------
p2 = paras[2]
set_text(p2,
    f"Task 2's modelling framework was executed as a baseline check (500 customers, 16.0% delinquent): "
    f"logistic regression, gradient boosting and a dummy baseline all scored at the no-skill floor "
    f"(PR-AUC {pr_auc_dm:.2f}-{pr_auc_gb:.2f} vs a {base_rate:.2f} base rate), confirming Task 1 - no "
    "recorded variable is yet a validated predictor. The segments below are the strongest observed "
    "(descriptive) gradients: fit for a controlled pilot, not automated decisions. Per the Dataset "
    "Description Guide, Missed_Payments covers 12 months and Month_1-6 cover 6 months; neither related "
    "to the target (p = 0.56).", size=9.5)

a = add_par_after(p2, "- Heavy debt burden: Debt_to_Income_Ratio above 40% delinquent at 18.9% (14/74) - "
                      "nearly double the 9.8% (8/82) below 20% (strongest gradient; not significant, p = 0.10).")
a = add_par_after(a, "- Income instability: unemployed delinquent at 19.4% (18/93) vs 11.5% (10/87) retired; "
                     "Employment_Status needed label clean-up ('EMP'/'employed', 245 rows).")
a = add_par_after(a, "- Product concentration: Business cardholders delinquent at 21.3% (23/108) - highest "
                     "of five Credit_Card_Type tiers - vs 11.8% (9/76) Platinum.")
a = add_par_after(a, "Top 3 Risk Factors (for easy communication):", bold=True)
a = add_par_after(a, "1. Debt-to-income ratio above 40% - roughly double the rate of the lowest band.",
                  style="List Bullet")
a = add_par_after(a, "2. Unemployed customers - the highest delinquency of any employment group (19.4%).",
                  style="List Bullet")
a = add_par_after(a, "3. Business card tier - the highest of all product tiers (21.3%).", style="List Bullet")
a = add_par_after(a, "All three are descriptive candidates; none is statistically validated (Task 2 "
                     "go/no-go gate applies).", size=9)

p3 = paras[3]
set_text(p3, "Key Insights Summary Table:", bold=True, size=9.5)

tbl = doc.tables[0]
rows_data = [
    ("Debt burden roughly doubles delinquency", "Accounts with DTI > 40% (n = 74, ~15% of book)",
     "Debt_to_Income_Ratio; Loan_Balance", "Priority cohort for proactive restructuring outreach"),
    ("Income instability elevates risk", "Unemployed customers (n = 93)",
     "Employment_Status; Income", "Route to hardship and collections-support programmes"),
    ("Risk concentrated in one product tier", "Business cardholders (n = 108)",
     "Credit_Card_Type", "Product-level review of limits and underwriting"),
]
for i, row in enumerate(rows_data, start=1):
    cells = tbl.rows[i].cells if i < len(tbl.rows) else tbl.add_row().cells
    for c, val in zip(cells, row):
        clear_runs(c.paragraphs[0])
        r = c.paragraphs[0].add_run(val); r.font.size = Pt(8)

# ---------------- 2. Recommendation Framework ----------------
set_text(paras[6],
    "The recommendation converts the strongest observed EDA insight into a SMART intervention, "
    "respecting the Task 2 finding that no model may yet be deployed.", size=9.5)
set_text(paras[7], "")

set_labeled(paras[8], "Restated Insight: ",
    "DTI above 40% delinquent at 18.9% vs 9.8% below 20% - the strongest observed gradient, not yet "
    "statistically confirmed.")
set_labeled(paras[9], "Proposed Recommendation: ",
    "Run a controlled 90-day Collections pilot: proactive, supportive outreach (hardship restructuring, "
    "limit review) to every account with DTI above 40%, against a matched control.")
set_labeled(paras[10], "Specific: ",
    "Target the DTI > 40% cohort (74 accounts, ~15% of book); matched control from the 20-30% band.")
set_labeled(paras[11], "Measurable: ",
    ">= 15% relative delinquency reduction vs control at day 90; track contact, cure and complaint "
    "rates and cost per account, monthly.")
set_labeled(paras[12], "Actionable: ",
    "Existing Collections team and CRM; rule-based campaign, no model deployment (Task 2 gate).")
set_labeled(paras[13], "Relevant: ",
    "Targets the highest-observed-delinquency segment, supporting credit-loss reduction while data "
    "remediation proceeds.")
set_labeled(paras[14], "Time-bound: ",
    "Launch within 4 weeks, run 90 days, scale-or-stop decision at day 91.")
set_labeled(paras[15], "Justification and Business Rationale: ",
    "Lowest-cost way to test the top observed gradient under controlled conditions: if confirmed, "
    "Geldium gains a validated, explainable targeting rule; if not, the gate redirects effort to data "
    "remediation instead of acting on noise. Transparent, feasible with current resources, and avoids "
    f"model investment while predictors remain unvalidated (PR-AUC at the {base_rate:.2f} no-skill floor).")

# ---------------- 3. Ethical and Responsible AI ----------------
set_text(paras[17],
    "With no validated signal in the data, the main ethical risk is acting on noise as if it were "
    "fact. Two fairness risks and mitigations:", size=9.5)
set_labeled(paras[18], "Fairness risk 1 - proxy discrimination: ",
    "Location and Employment_Status correlate with ethnicity and socioeconomic status; punitive use "
    "(limit cuts, rejections) could disadvantage protected groups. Mitigation: DTI targeting for "
    "supportive offers only, exclude Location from scoring, audit pilot outcomes by age, employment "
    "group and region.")
set_labeled(paras[19], "Fairness risk 2 - vulnerable-group impact: ",
    "Unemployed customers already show the highest delinquency (19.4%); aggressive collections could "
    "deepen hardship. Mitigation: human-in-the-loop review, hardship-programme routing, plain-language "
    "reason codes (e.g., 'your debt obligations are close to your income').")
set_labeled(paras[20], "Explainability for non-technical stakeholders: ",
    "The recommendation is a transparent DTI-band rule, not a black box; any future model must give "
    "per-decision reason codes. Predictions are communicated as probabilities with uncertainty - the "
    "executed baseline showed no better than chance, so no automated adverse decisions are justified.")
set_labeled(paras[21], "Responsible and transparent AI use: ",
    "Transparency (pilot limits documented, including 7.8% missing Income affecting banding), "
    "accountability (named owner, model card, stop rule), data privacy (income/debt data minimised), "
    "and a fairness re-audit before scale-up.")

set_text(paras[22], "")

# --- tighten spacing on body paragraphs ---
for p in doc.paragraphs:
    if p.style.name in ("Normal", "List Bullet", "List Bullet 2"):
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

doc.core_properties.title = "Geldium Business Summary Report - Predictive Insights for Collections Strategy"
doc.core_properties.author = "Tata iQ Forage Simulation - Data Analytics"
doc.save(OUT)
print("saved:", OUT)

# %% [markdown]
# ## QA

# %%
d2 = docx.Document(OUT)
print("headings:", [p.text for p in d2.paragraphs if p.style.name.startswith("Heading")])
print("tables:", len(d2.tables), "| table rows:", len(d2.tables[0].rows))
left = [p.text[:60] for p in d2.paragraphs if "[Insert" in p.text or "Describe the" in p.text]
print("placeholders left:", left or "none")
