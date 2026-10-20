# %% [markdown]
# # Task 4 — AI-Powered Collections Strategy deck (Tata iQ Forage Simulation)
#
# Fills the official `Presentation_Template.pptx` (marina master, dark-blue theme, Arial) with
# evidence-grounded content from Tasks 1-3. Design follows the template theme: white text on the
# #00517C background, accent-coloured chevrons/tables, Tata logo badge.
#
# Input : ../Presentation_Template.pptx (preserved, never modified)
# Output: AI_Powered_Collections_Strategy_Deck.pptx

# %%
import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BASE = Path(__file__).resolve().parent
ROOT = BASE.parent
TEMPLATE = ROOT / "Presentation_Template.pptx"
OUT = BASE / "AI_Powered_Collections_Strategy_Deck.pptx"
LOGO = ROOT / "Task_1_EDA" / "assets" / "tata_logo.png"
CHART = ROOT / "Task_1_EDA" / "Geldium_EDA_Outputs" / "Geldium_EDA_Charts" / "fig4_delinq_by_dti.png"

# theme palette (from template)
NAVY_BG   = RGBColor(0x00, 0x51, 0x7C)
ACCENT    = RGBColor(0x02, 0x77, 0xBD)   # accent1 blue
TEAL      = RGBColor(0x00, 0x96, 0x88)   # accent3
SKY       = RGBColor(0x03, 0x9B, 0xE5)   # accent4
GREEN     = RGBColor(0x55, 0x8B, 0x2F)   # accent2
LIGHT     = RGBColor(0xCF, 0xD8, 0xDC)   # lt2
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
DARKTXT   = RGBColor(0x00, 0x40, 0x65)   # dk2
YELLOW    = RGBColor(0xFF, 0xEB, 0x38)   # accent6

prs = Presentation(TEMPLATE)
S = prs.slides

def clear_tf(tf):
    tf.clear()
    return tf

def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def add_textbox(slide, l, t, w, h, lines, size=10, color=WHITE, bold_first=False,
                align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=4):
    """lines: list of str OR (bold_lead, rest) tuples."""
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if isinstance(line, tuple):
            r1 = p.add_run(); r1.text = line[0]
            r1.font.bold = True; r1.font.size = Pt(size); r1.font.name = "Arial"
            r1.font.color.rgb = YELLOW
            r2 = p.add_run(); r2.text = line[1]
            r2.font.size = Pt(size); r2.font.name = "Arial"; r2.font.color.rgb = color
        else:
            r = p.add_run(); r.text = line
            r.font.size = Pt(size); r.font.name = "Arial"; r.font.color.rgb = color
            r.font.bold = bold_first and i == 0
    return tb

def style_table(tbl, header_fill=ACCENT, body_fill=LIGHT, header_size=10.5, body_size=9.5):
    for ri, row in enumerate(tbl.rows):
        for cell in row.cells:
            cell.margin_left = Inches(0.06); cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if ri == 0 else body_fill
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(header_size if ri == 0 else body_size)
                    r.font.bold = (ri == 0)
                    r.font.color.rgb = WHITE if ri == 0 else DARKTXT

# =====================================================================
# SLIDE 1 — Title
# =====================================================================
s1 = S[0]
# keep official title & subtitle; add logo badge + byline
badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.85), Inches(0.30),
                            Inches(0.90), Inches(0.90))
badge.fill.solid(); badge.fill.fore_color.rgb = WHITE; badge.line.fill.background()
s1.shapes.add_picture(str(LOGO), Inches(9.00), Inches(0.38), width=Inches(0.60))
add_textbox(s1, 1.84, 4.45, 6.32, 0.5,
            ["Tata iQ Forage Simulation  |  Predictive Analytics Team  |  October 2026"],
            size=11, color=LIGHT, align=PP_ALIGN.LEFT)
set_notes(s1, "Opening frame: this deck presents an agentic-AI collections system for Geldium. "
              "Every claim is grounded in the executed Task 1 EDA (500 customers, 16.0% delinquent) "
              "and the executed Task 2 baseline check - we deliberately pilot before we model.")

# =====================================================================
# SLIDE 2 — How the System Works (4-part flow diagram)
# =====================================================================
s2 = S[1]
clear_tf(s2.shapes[1].text_frame)  # body placeholder -> blank, we draw our own

steps = [
    ("1 · INPUTS", ACCENT, [
        "Customer data: Income, Credit_Score, Credit_Utilization, Debt-to-Income, Loan_Balance, Missed_Payments (12 mo), 6-month payment history, Employment_Status, Tenure",
        "Data-quality gate first: label clean-up, median imputation + missingness flags (Task 1 cleaning log)",
    ]),
    ("2 · DECISION LOGIC", TEAL, [
        "Rules first: transparent DTI-band risk tiers (low / medium / high)",
        "Predictive model only after the go/no-go gate (cross-validated PR-AUC must beat the 0.16 no-skill floor)",
        "Calibrated probability -> risk score + reason codes",
    ]),
    ("3 · ACTIONS", SKY, [
        "Low risk: automated reminder e-mails / SMS",
        "Medium risk: proactive hardship or limit-review offer",
        "High risk or vulnerable: human collections review",
    ]),
    ("4 · LEARNING LOOP", GREEN, [
        "Outcomes tracked: cure rate, roll-rate, complaints, cost per account",
        "Monthly fairness + drift (PSI) audits",
        "Rules and models refreshed; kill-switch to manual review",
    ]),
]
x0, box_w, gap = 0.30, 2.30, 0.05
for i, (label, color, detail) in enumerate(steps):
    x = x0 + i * (box_w + gap)
    chev = s2.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(1.20),
                               Inches(box_w), Inches(0.55))
    chev.fill.solid(); chev.fill.fore_color.rgb = color
    chev.line.color.rgb = WHITE; chev.line.width = Pt(0.75)
    tf = chev.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.bold = True; r.font.size = Pt(11.5); r.font.name = "Arial"; r.font.color.rgb = WHITE
    card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.85),
                               Inches(box_w), Inches(2.55))
    card.fill.solid(); card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = color; card.line.width = Pt(1.25)
    tf = card.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.08); tf.margin_bottom = Inches(0.06)
    for j, d in enumerate(detail):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        r = p.add_run(); r.text = "- " + d
        r.font.size = Pt(8.5); r.font.name = "Arial"; r.font.color.rgb = DARKTXT
add_textbox(s2, 0.30, 4.60, 9.4, 0.4,
            ["Baseline: 500 customers, 16.0% delinquent - every pipeline step is documented in the "
             "Task 1 cleaning log and Task 2 model card."],
            size=9, color=LIGHT)
set_notes(s2, "Walk left to right. Emphasise the data-quality gate: Task 1 found 6 label variants in "
              "Employment_Status (245 rows fixed), 7.8% missing Income (median imputed + flagged) and "
              "81 tenure-age inconsistencies - the system consumes only cleaned, documented data. "
              "Decision logic is rules-first because the executed baseline showed models at the "
              "no-skill floor (PR-AUC 0.15-0.17 vs 0.16 base rate).")

# =====================================================================
# SLIDE 3 — Role of Agentic AI (autonomous vs human oversight table)
# =====================================================================
s3 = S[2]
clear_tf(s3.shapes[1].text_frame)
rows = [
    ("AUTONOMOUS  (AI acts, no approval needed)", "HUMAN OVERSIGHT  (approval required)"),
    ("Nightly segmentation and risk scoring (rules + validated model)",
     "Model go/no-go approval and threshold setting (Task 2 gate)"),
    ("Reminder e-mails / SMS to low-risk accounts",
     "Hardship and restructuring offers (affordability judgement)"),
    ("Case routing and priority queues for collectors",
     "Any adverse action: limit cuts, escalation, legal action"),
    ("Drift / anomaly detection and automatic alerting",
     "Fairness and bias review before any scale-up"),
    ("Outcome dashboards and standard reporting",
     "Vulnerable-customer exceptions and complaint handling"),
]
tbl_shape = s3.shapes.add_table(len(rows), 2, Inches(0.42), Inches(1.35),
                                Inches(9.15), Inches(3.30))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(4.57); tbl.columns[1].width = Inches(4.58)
for ri, (a, b) in enumerate(rows):
    tbl.cell(ri, 0).text = a; tbl.cell(ri, 1).text = b
style_table(tbl)
# highlight header of right column in teal for contrast
tbl.cell(0, 1).fill.solid(); tbl.cell(0, 1).fill.fore_color.rgb = TEAL
add_textbox(s3, 0.42, 4.75, 9.15, 0.35,
            ["Principle: AI acts autonomously only where actions are reversible, low-impact and fully "
             "explainable; everything affecting a customer's credit terms stays human."],
            size=9, color=LIGHT)
set_notes(s3, "Agentic AI here means the system can plan and execute routine steps end-to-end (scoring, "
              "messaging, routing, monitoring) without ticket queues. The boundary is deliberate: "
              "unemployed customers show the highest observed delinquency (19.4%), so vulnerable or "
              "adverse-touching decisions always route to a human officer.")

# =====================================================================
# SLIDE 4 — Responsible AI Guardrails
# =====================================================================
s4 = S[3]
clear_tf(s4.shapes[1].text_frame)
add_textbox(s4, 0.42, 1.30, 9.15, 3.6, [
    ("Fairness by design - ",
     "subgroup audits (Age band, Employment_Status, Location) for equal-opportunity and calibration "
     "parity; Location excluded from scoring because it proxies ethnicity and socio-economic status."),
    ("Explainability - ",
     "transparent DTI-band rules today; per-decision reason codes for any future model; predictions "
     "communicated as probabilities with uncertainty (executed baseline: no better than chance)."),
    ("Human-in-the-loop and compliance - ",
     "no fully automated adverse decisions; adverse-action reason notices; income and debt data "
     "minimised to need-to-know; documented model card and full audit trail."),
    ("Continuous monitoring - ",
     "monthly drift (PSI), calibration and complaint-rate reviews; kill-switch back to manual review; "
     "fairness re-audit before any scale-up."),
], size=12.5, space_after=14)
add_textbox(s4, 0.42, 4.75, 9.15, 0.35,
            ["Guardrails are build requirements, not afterthoughts - each maps to a Task 2 evaluation "
             "metric and a named owner."],
            size=9, color=LIGHT)
set_notes(s4, "Two concrete fairness risks drive these guardrails: (1) proxy discrimination via "
              "Location / Employment_Status, (2) vulnerable-group impact on unemployed customers "
              "(19.4% delinquent). Compliance angle: adverse-action transparency and data minimisation "
              "for income/debt fields; 7.8% missing Income is disclosed as a banding limitation.")

# =====================================================================
# SLIDE 5 — Expected Business Impact (KPIs vs customer outcomes)
# =====================================================================
s5 = S[4]
clear_tf(s5.shapes[1].text_frame)
rows = [
    ("BUSINESS KPIs", "CUSTOMER OUTCOMES"),
    ("15%+ relative delinquency reduction in the pilot cohort (DTI > 40%, 74 accounts) by day 90",
     "Supportive offers before punitive action; hardship routing for vulnerable customers"),
    ("Lower cost per account: automated reminders replace manual dialling",
     "Fewer intrusive contacts for low-risk customers"),
    ("Higher cure rate and lower 90+ DPD roll-rate in the targeted band",
     "Transparent reason codes (e.g., 'your debts are close to your income')"),
    ("Collections capacity freed for complex, high-value cases",
     "Fair outcomes audited by age, employment group and region"),
    ("Near-zero marginal cost as the portfolio scales",
     "Consistent, documented process customers can question and challenge"),
]
tbl_shape = s5.shapes.add_table(len(rows), 2, Inches(0.42), Inches(1.30),
                                Inches(9.15), Inches(3.35))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(4.57); tbl.columns[1].width = Inches(4.58)
for ri, (a, b) in enumerate(rows):
    tbl.cell(ri, 0).text = a; tbl.cell(ri, 1).text = b
style_table(tbl, header_fill=GREEN)
tbl.cell(0, 1).fill.solid(); tbl.cell(0, 1).fill.fore_color.rgb = ACCENT
add_textbox(s5, 0.42, 4.75, 9.15, 0.35,
            ["Impact targets come from the Task 3 SMART pilot - measured against a matched control, "
             "with a scale-or-stop decision at day 91."],
            size=9, color=LIGHT)
set_notes(s5, "Quantitative targets are pilot commitments, not modelled projections: 15% relative "
              "delinquency reduction in the DTI > 40% cohort (18.9% observed vs 9.8% in the lowest "
              "band). Qualitative gains - fairness, transparency, scalability - are audited with the "
              "same rigor as the KPIs.")

# =====================================================================
# SLIDE 6 — Evidence base (why pilot-first)
# =====================================================================
s6 = S[5]
title6 = s6.shapes[0]
title6.text_frame.paragraphs[0].runs[0].text = "Evidence Base: Why Pilot-First"
clear_tf(s6.shapes[1].text_frame)
card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.42), Inches(1.30),
                           Inches(5.05), Inches(3.30))
card.fill.solid(); card.fill.fore_color.rgb = WHITE
card.line.color.rgb = ACCENT; card.line.width = Pt(1.25)
s6.shapes.add_picture(str(CHART), Inches(0.55), Inches(1.42), width=Inches(4.80))
add_textbox(s6, 5.70, 1.30, 3.90, 3.4, [
    ("Base rate: ", "16.0% delinquent (80 of 500 customers)"),
    ("Strongest observed gradient: ", "DTI > 40% delinquent at 18.9% vs 9.8% below 20% (~2x)"),
    ("Executed baseline: ", "logistic regression, boosting and dummy all at PR-AUC 0.15-0.17, "
                            "i.e. the 0.16 no-skill floor"),
    ("Implication: ", "no model deployment yet - validate the DTI gradient with a controlled "
                      "90-day pilot, then scale what works"),
    ("Definitions: ", "cross-checked against the official Dataset Description Guide "
                      "(Missed_Payments = 12-month window; Month_1-6 = 6-month window)"),
], size=10.5, space_after=8)
add_textbox(s6, 0.42, 4.75, 9.15, 0.35,
            ["Caveat: segments are descriptive, not statistically validated - the pilot exists to "
             "validate before scaling (Task 2 go/no-go gate)."],
            size=9, color=YELLOW)
set_notes(s6, "This slide keeps us honest: the same EDA that found the DTI gradient found that no "
              "predictor beats chance today. Pilot-first is how we capture upside without acting on "
              "noise - and every number here is reproducible from the project repository.")

prs.save(OUT)
print("saved:", OUT)

# %% [markdown]
# ## QA

# %%
chk = Presentation(OUT)
print("slides:", len(chk.slides))
for i, s in enumerate(chk.slides, 1):
    title = ""
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            title = sh.text_frame.text.strip().splitlines()[0][:60]
            break
    n_shapes = len(s.shapes)
    has_notes = bool(s.notes_slide.notes_text_frame.text.strip())
    print(f"  slide {i}: shapes={n_shapes} notes={'yes' if has_notes else 'NO'} first_text={title!r}")
